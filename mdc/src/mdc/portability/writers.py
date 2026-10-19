"""Exporters: a database -> downloadable file bytes."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any

from mdc.databases.manager import DatabaseHandle
from mdc.model.operation import ReadOperation
from mdc.portability.plan import RECORD_ID, ExportOptionError
from mdc.portability.readers import EXPORT_FORMAT_NAME, EXPORT_FORMAT_VERSION

# format -> (default scope, allowed scopes)
FORMATS: dict[str, tuple[str, tuple[str, ...]]] = {
    "dnadb": ("all", ("schema", "data", "all")),
    "sql": ("all", ("schema", "data", "all")),
    "xlsx": ("all", ("schema", "data", "all")),
    "csv": ("data", ("data", "all")),
    "md": ("schema", ("schema",)),
    "er": ("schema", ("schema",)),
}
FORMAT_ALIASES = {"xls": "xlsx", "excel": "xlsx", "json": "dnadb", "markdown": "md", "erd": "er", "mermaid": "er", "script": "sql"}
SCOPES = ("schema", "data", "all")

_MEDIA = {
    "dnadb": "application/json", "sql": "application/sql", "csv": "text/csv", "md": "text/markdown",
    "er": "text/plain", "zip": "application/zip",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_SQL_TYPES = {"string": "VARCHAR", "decimal": "DOUBLE PRECISION", "integer": "BIGINT", "boolean": "BOOLEAN", "datetime": "TIMESTAMP"}


@dataclass(frozen=True)
class ExportFile:
    filename: str
    media_type: str
    content: bytes


def resolve_export_options(fmt: str | None, scope: str | None, *, table: str | None = None) -> tuple[str, str]:
    fmt = FORMAT_ALIASES.get((fmt or "dnadb").lower(), (fmt or "dnadb").lower())
    if fmt not in FORMATS:
        raise ExportOptionError(f"Unknown export format {fmt!r}. Choose one of: csv, sql, xlsx, md, er, dnadb.")
    default_scope, allowed = FORMATS[fmt]
    scope = (scope or default_scope).lower()
    if scope not in SCOPES:
        raise ExportOptionError(f"Unknown scope {scope!r}. Choose one of: schema, data, all.")
    if scope not in allowed:
        reason = {
            "csv": "CSV holds data only - for structure use md, sql, xlsx or dnadb.",
            "md": "A design document is structure only.",
            "er": "An ER diagram is structure only.",
        }.get(fmt, "")
        raise ExportOptionError(f"{fmt} can't be exported with scope {scope!r}. {reason}".strip())
    if table is not None and fmt not in ("csv", "sql", "dnadb", "xlsx", "md", "er"):
        raise ExportOptionError("table is not supported for this format.")
    return fmt, "data" if (fmt == "csv" and scope == "all") else scope


def export_filename(database: str, fmt: str, scope: str, table: str | None = None) -> str:
    base = f"{database}-{table}" if table else database
    suffix = "" if scope == "all" or fmt in ("csv", "md", "er") else f"-{'structure' if scope == 'schema' else 'data'}"
    if fmt == "csv":
        # Named after the table alone: importing a CSV takes its table name
        # from the filename, so "<db>-<table>.csv" would come back as a table
        # called "<db>_<table>".
        return f"{table}.csv" if table else f"{database}-csv.zip"
    return {
        "dnadb": f"{base}{suffix}.dnadb.json", "sql": f"{base}{suffix}.sql", "xlsx": f"{base}{suffix}.xlsx",
        "md": f"{base}-design.md", "er": f"{base}-er.mmd",
    }[fmt]


def export_database(handle: DatabaseHandle, fmt: str | None, scope: str | None = None, table: str | None = None) -> ExportFile:
    fmt, scope = resolve_export_options(fmt, scope, table=table)
    tables = _tables(handle, table)
    now = datetime.now(timezone.utc)
    name = export_filename(handle.name, fmt, scope, table)
    if fmt == "dnadb":
        return ExportFile(name, _MEDIA["dnadb"], _native(handle, tables, scope, now))
    if fmt == "sql":
        return ExportFile(name, _MEDIA["sql"], _sql(handle, tables, scope, now).encode("utf-8"))
    if fmt == "xlsx":
        return ExportFile(name, _MEDIA["xlsx"], _xlsx(handle, tables, scope))
    if fmt == "md":
        return ExportFile(name, _MEDIA["md"], _markdown(handle.name, tables, now).encode("utf-8"))
    if fmt == "er":
        return ExportFile(name, _MEDIA["er"], _mermaid(tables).encode("utf-8"))
    if table:
        return ExportFile(name, _MEDIA["csv"], _csv_text(handle, table, tables[table]).encode("utf-8"))
    return ExportFile(name, _MEDIA["zip"], _csv_zip(handle, tables))


# -- reading a database ---------------------------------------------------------------

Fields = dict[str, Any]  # field name -> FieldSchema


def _tables(handle: DatabaseHandle, only: str | None) -> dict[str, Fields]:
    names = handle.schema_registry.list_collections()
    if only is not None:
        if only not in names:
            from mdc.schema.registry import CollectionNotFoundError
            raise CollectionNotFoundError(only)
        names = [only]
    return {n: dict(handle.schema_registry.get_collection(n).fields) for n in names}


def _rows(handle: DatabaseHandle, table: str) -> list[dict[str, Any]]:
    records = handle.engine.read(ReadOperation(collection=table, filters=[])).records
    records = sorted(records, key=lambda r: (r.created_at, r.record_id))
    return [{RECORD_ID: r.record_id, **r.fields} for r in records]


def _plain(value: Any) -> Any:
    return value.isoformat() if isinstance(value, datetime) else value


# -- MDC export file ------------------------------------------------------------------

def _native(handle: DatabaseHandle, tables: dict[str, Fields], scope: str, now: datetime) -> bytes:
    body: dict[str, Any] = {}
    for name, fields in tables.items():
        entry: dict[str, Any] = {}
        if scope in ("schema", "all"):
            entry["fields"] = {f.name: {"datatype": f.datatype, "required": f.required} for f in fields.values()}
        if scope in ("data", "all"):
            entry["rows"] = [
                {"record_id": row[RECORD_ID], "fields": {k: _plain(v) for k, v in row.items() if k != RECORD_ID}}
                for row in _rows(handle, name)
            ]
        body[name] = entry
    bundle = {
        "format": EXPORT_FORMAT_NAME,
        "version": EXPORT_FORMAT_VERSION,
        "exported_at": now.isoformat(),
        "source_database": handle.name,
        "includes": {"structure": scope in ("schema", "all"), "data": scope in ("data", "all")},
        "tables": body,
    }
    return json.dumps(bundle, indent=2).encode("utf-8")


# -- CSV ------------------------------------------------------------------------------

def _guard_cell(value: str) -> str:
    """A spreadsheet runs a cell starting with = + - @ as a formula. A
    leading ' makes it text; `readers.unguard_cell` removes it again on
    import, so the round trip stays lossless."""
    return "'" + value if value[:1] in ("=", "+", "-", "@") else value


def _csv_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return _guard_cell(str(_plain(value))) if isinstance(_plain(value), str) else str(value)


def _csv_text(handle: DatabaseHandle, table: str, fields: Fields) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    columns = [RECORD_ID, *fields]
    writer.writerow(columns)
    for row in _rows(handle, table):
        writer.writerow([_csv_cell(row.get(c)) for c in columns])
    return buffer.getvalue()


def _csv_zip(handle: DatabaseHandle, tables: dict[str, Fields]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, fields in tables.items():
            archive.writestr(f"{name}.csv", _csv_text(handle, name, fields))
    return buffer.getvalue()


# -- SQL ------------------------------------------------------------------------------

def _q(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def _sql_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return repr(value)
    return "'" + str(_plain(value)).replace("'", "''") + "'"


def _sql(handle: DatabaseHandle, tables: dict[str, Fields], scope: str, now: datetime) -> str:
    out = [
        f'-- MDC export of database "{handle.name}" - {now.strftime("%Y-%m-%d %H:%M UTC")}',
        f"-- mdc-database: {handle.name}",
        f"-- {'Structure and data' if scope == 'all' else 'Structure only' if scope == 'schema' else 'Data only'}. "
        "Standard SQL: import it into MDC, or run it in another database.",
        "",
    ]
    if scope in ("schema", "all"):
        for name, fields in tables.items():
            cols = [f"  {_q(RECORD_ID)} VARCHAR(64) PRIMARY KEY"]
            cols += [f"  {_q(f.name)} {_SQL_TYPES[f.datatype]}{' NOT NULL' if f.required else ''}" for f in fields.values()]
            out.append(f"CREATE TABLE {_q(name)} (\n" + ",\n".join(cols) + "\n);")
            out.append("")
    if scope in ("data", "all"):
        for name, fields in tables.items():
            rows = _rows(handle, name)
            columns = [RECORD_ID, *fields]
            head = f"INSERT INTO {_q(name)} ({', '.join(_q(c) for c in columns)}) VALUES"
            for start in range(0, len(rows), 100):
                chunk = rows[start : start + 100]
                values = ",\n".join("  (" + ", ".join(_sql_literal(r.get(c)) for c in columns) + ")" for r in chunk)
                out.append(f"{head}\n{values};")
            if rows:
                out.append("")
    return "\n".join(out)


# -- Excel ----------------------------------------------------------------------------

def _sheet_title(name: str, used: set[str]) -> str:
    title = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "sheet"
    candidate, n = title, 2
    while candidate.lower() in used:
        suffix = f"~{n}"
        candidate = title[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _xlsx(handle: DatabaseHandle, tables: dict[str, Fields], scope: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = {"_schema"}

    def add_row(ws, values: list[Any], bold: bool = False) -> None:
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # text, never a formula, even if it starts with "="
            if bold:
                cell.font = Font(bold=True)

    if scope in ("data", "all"):
        for name, fields in tables.items():
            ws = wb.create_sheet(_sheet_title(name, used))
            columns = [RECORD_ID, *fields]
            add_row(ws, columns, bold=True)
            for row in _rows(handle, name):
                add_row(ws, [_plain(row.get(c)) for c in columns])
            ws.freeze_panes = "A2"
    if scope in ("schema", "all"):
        ws = wb.create_sheet("_schema")
        add_row(ws, ["table", "field", "type", "required"], bold=True)
        for name, fields in tables.items():
            for f in fields.values():
                add_row(ws, [name, f.name, f.datatype, "yes" if f.required else "no"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# -- Markdown design document and ER diagram ------------------------------------------

def _relationships(tables: dict[str, Fields]) -> list[tuple[str, str, str]]:
    """(child table, parent table, field). MDC stores no foreign keys, so
    these are inferred from a `<table>_id` field naming convention."""
    found: list[tuple[str, str, str]] = []
    for child, fields in tables.items():
        for field_name in fields:
            if not field_name.endswith("_id") or field_name == RECORD_ID:
                continue
            base = field_name[:-3]
            for candidate in (base, base + "s", base + "es", base[:-1] + "ies" if base.endswith("y") else None):
                if candidate and candidate in tables and candidate != child:
                    found.append((child, candidate, field_name))
                    break
    return found


def _mermaid_name(name: str) -> str:
    return name if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name) else '"' + name.replace('"', "'") + '"'


def _mermaid(tables: dict[str, Fields]) -> str:
    lines = ["erDiagram"]
    for name, fields in tables.items():
        lines.append(f"    {_mermaid_name(name)} {{")
        lines.append(f"        string {RECORD_ID} PK")
        for f in fields.values():
            attr = re.sub(r"[^A-Za-z0-9_]", "_", f.name)
            lines.append(f"        {f.datatype} {attr}" + (' "required"' if f.required else ""))
        lines.append("    }")
    for child, parent, via in _relationships(tables):
        lines.append(f'    {_mermaid_name(child)} }}o--|| {_mermaid_name(parent)} : "{via}"')
    return "\n".join(lines) + "\n"


def _markdown(database: str, tables: dict[str, Fields], now: datetime) -> str:
    out = [
        f"# Database: {database}",
        "",
        f"_MDC design document - exported {now.strftime('%Y-%m-%d')} - {len(tables)} table(s) - structure only_",
        "",
        "## Entity-relationship diagram",
        "",
        "```mermaid",
        _mermaid(tables).rstrip(),
        "```",
        "",
    ]
    inferred = _relationships(tables)
    out.append(
        "Relationships are inferred from `<table>_id` field names - MDC does not store foreign keys."
        if inferred else "No relationships are drawn: MDC does not store foreign keys, and no field is named `<table>_id`."
    )
    out += ["", f"Every row also carries a system-generated `{RECORD_ID}` (not listed below; it is not an importable field).", ""]
    for name, fields in tables.items():
        out += [f"## Table: {name}", "", "| Field | Type | Required |", "|---|---|---|"]
        out += [f"| {f.name} | {f.datatype} | {'yes' if f.required else 'no'} |" for f in fields.values()]
        out.append("")
    return "\n".join(out)
